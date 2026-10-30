"""備品棚卸: 保存済みデータのエクセル出力（/api/kintai/bihin/xlsx）。"""
from __future__ import annotations

import hashlib
import json
import re
import sqlite3
import subprocess
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import kintai
import kintai_bihin
import kintai_bihin_xlsx

ROOT = Path(__file__).resolve().parents[1]
JS = (ROOT / 'kintai-bihin.js').read_text(encoding='utf-8')
HTML = (ROOT / 'index.html').read_text(encoding='utf-8')
NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'


def read_xlsx(blob: bytes):
    """stdlib だけで .xlsx を開き {ref: (value, type, numFmt)} を返す。"""
    with ZipFile(BytesIO(blob)) as z:
        sheet = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        styles = ET.fromstring(z.read('xl/styles.xml'))
        names = z.namelist()
    custom = {int(n.get('numFmtId')): n.get('formatCode') for n in styles.iter(NS + 'numFmt')}
    builtin = {0: 'General', 3: '#,##0'}
    xfs = [int(x.get('numFmtId')) for x in styles.find(NS + 'cellXfs')]
    cells = {}
    for c in sheet.iter(NS + 'c'):
        fmt = xfs[int(c.get('s', 0))]
        fmt = custom.get(fmt, builtin.get(fmt))
        if c.get('t') == 'inlineStr':
            cells[c.get('r')] = (''.join(t.text or '' for t in c.iter(NS + 't')), 's', fmt)
        elif c.find(NS + 'v') is not None:
            v = c.find(NS + 'v').text
            cells[c.get('r')] = (float(v) if '.' in v else int(v), 'n', fmt)
        else:
            cells[c.get('r')] = (None, 'blank', fmt)
    merges = [m.get('ref') for m in sheet.iter(NS + 'mergeCell')]
    return cells, merges, names


class _H:
    def __init__(self, user, path):
        self.user = user
        self.path = path


class XlsxExportTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.c = sqlite3.connect(Path(self.tmp.name) / 't.sqlite3')
        self.c.row_factory = sqlite3.Row
        self.c.execute(
            "CREATE TABLE users (id INTEGER PRIMARY KEY, login TEXT, name TEXT, password TEXT, "
            "role TEXT, manage INTEGER, overview INTEGER, active INTEGER, must_change INTEGER, "
            "version INTEGER, mark_kintai INTEGER DEFAULT 0)"
        )
        kintai.bind(Path(self.tmp.name))
        kintai.schema(self.c)
        self.user = {'id': 1, 'name': '主', 'login': 'o@ex.com', 'role': 'owner', 'mark_kintai': 1}
        for cat, n, unit, p in (('文具', 'ボールペン', '本', 120), ('資材', 'OSKラベル', '巻', 38000), ('資材', '宅配袋', '枚', 32.5)):
            kintai_bihin.save_item(self.c, self.user, {'category': cat, 'name': n, 'unit': unit, 'unitPrice': p})
        rows = kintai_bihin.get_sheet(self.c, 2026, 10)['rows']
        by = {r['name']: r for r in rows}
        # 保存順は マスタ順と逆（宅配袋→ボールペン→OSK）
        saved = [
            dict(by['宅配袋'], openingQty=3, purchaseQty=0, closingQty=2),   # 3×32.5=97.5→98, 購入0は0, 2×32.5=65, 使用1→32.5→33
            dict(by['ボールペン'], openingQty=10, purchaseQty=5, closingQty=None),  # 月末空→使用も空
            dict(by['OSKラベル'], openingQty=None, purchaseQty=None, closingQty=None),
        ]
        code, _ = kintai_bihin.save_sheet(self.c, self.user, {'year': 2026, 'month': 10, 'version': 0, 'rows': saved})
        self.assertEqual(code, 200)
        self.c.commit()

    def tearDown(self):
        self.c.close()
        self.tmp.cleanup()

    def _dump(self):
        return hashlib.sha256(''.join(self.c.iterdump()).encode()).hexdigest()

    def _get(self, user, query='?year=2026&month=10'):
        return kintai.handle_get(_H(user, '/api/kintai/bihin/xlsx' + query), self.c, '/api/kintai/bihin/xlsx')

    def test_values_blanks_totals_order(self):
        before = self._dump()
        res = self._get(self.user)
        self.assertEqual(self._dump(), before)  # 読み取り専用
        self.assertEqual(res[0], 200)
        self.assertEqual(res[2], kintai_bihin_xlsx.XLSX)
        cd = res[3]['Content-Disposition']
        cd.encode('latin-1')  # HTTP ヘッダーとして送れる
        m = re.search(r"filename\*=UTF-8''(\S+)", cd)
        self.assertEqual(unquote(m.group(1)), '備品棚卸_2026-10.xlsx')
        self.assertIn('filename="', cd)
        cells, merges, names = read_xlsx(res[1])
        self.assertIn('xl/workbook.xml', names)
        self.assertEqual(cells['A1'][0], '備品棚卸　2026年10月')
        self.assertTrue(cells['A2'][0].startswith('最終保存：'))
        self.assertIn('（主）', cells['A2'][0])
        heads = [cells[f'{c}3'][0] for c in 'ABCDEGIK']
        self.assertEqual(heads, ['科目', '品名', '単位', '単価', '月初在庫', '当月購入', '月末在庫', '当月使用'])
        self.assertEqual([cells[f'{c}4'][0] for c in 'EFGHIJKL'], ['数量', '金額'] * 4)
        # 保存順
        self.assertEqual([cells[f'B{r}'][0] for r in (5, 6, 7)], ['宅配袋', 'ボールペン', 'OSKラベル'])
        # 宅配袋（単価 32.5, 小数は四捨五入）
        self.assertEqual(cells['A5'][0], '資材')
        self.assertEqual(cells['D5'][:2], (32.5, 'n'))
        self.assertEqual(cells['E5'][:2], (3, 'n'))
        self.assertEqual(cells['F5'][:2], (98, 'n'))
        self.assertEqual(cells['G5'][:2], (0, 'n'))  # 0 は 0、空欄とは区別
        self.assertEqual(cells['H5'][:2], (0, 'n'))
        self.assertEqual(cells['I5'][0], 2)
        self.assertEqual(cells['J5'][0], 65)
        self.assertEqual(cells['K5'][0], 1)
        self.assertEqual(cells['L5'][0], 33)
        # ボールペン: 月末空 → 月末・使用は空欄（0 ではない）
        self.assertEqual([cells[f'{c}6'][0] for c in 'DEFGH'], [120, 10, 1200, 5, 600])
        for c in 'IJKL':
            self.assertEqual(cells[f'{c}6'][1], 'blank', c)
        # OSK: 数量すべて空
        for c in 'EFGHIJKL':
            self.assertEqual(cells[f'{c}7'][1], 'blank', c)
        self.assertEqual(cells['D7'][0], 38000)
        # 金額合計（画面・API と同じ値）
        totals = kintai_bihin.get_sheet(self.c, 2026, 10)['totals']
        self.assertEqual(cells['A8'][0], '金額合計')
        self.assertEqual(cells['F8'][0], 98 + 1200)
        self.assertEqual(cells['H8'][0], 600)
        self.assertEqual(cells['J8'][0], 65)
        self.assertEqual(cells['L8'][0], 33)
        self.assertEqual([cells[f'{c}8'][0] for c in 'FHJL'],
                         [totals['openingAmount'], totals['purchaseAmount'], totals['closingAmount'], totals['usageAmount']])
        self.assertIn('A8:D8', merges)
        # 書式: 金額・整数は #,##0 の数値セル
        for ref in ('D6', 'E6', 'F6', 'F5', 'F8', 'L8', 'D7'):
            self.assertEqual(cells[ref][1], 'n', ref)
            self.assertEqual(cells[ref][2], '#,##0', ref)

    def test_unsaved_month_is_rejected_and_untouched(self):
        before = self._dump()
        res = self._get(self.user, '?year=2026&month=11')
        self.assertEqual(res[0], 404)
        self.assertIn('保存されていない', res[1]['error'])
        self.assertEqual(self._dump(), before)
        self.assertIsNone(self.c.execute('SELECT 1 FROM kintai_bihin_sheets WHERE month=11').fetchone())

    def test_permission_matches_bihin_get(self):
        cases = [
            {'id': 2, 'name': '営', 'role': 'sales', 'manage': 0, 'overview': 0, 'mark_kintai': 0},
            {'id': 3, 'name': '現', 'role': 'floor3', 'manage': 0, 'overview': 0, 'mark_kintai': 0},
            {'id': 4, 'name': 'マ', 'role': 'sales', 'manage': 0, 'overview': 0, 'mark_kintai': 1},
            None,
        ]
        for u in cases:
            g = kintai.handle_get(_H(u, '/api/kintai/bihin?year=2026&month=10'), self.c, '/api/kintai/bihin')
            x = self._get(u)
            self.assertEqual(g[0], x[0], u)
            if g[0] == 403:
                self.assertEqual(g[1], x[1])

    def test_openpyxl_reads_numbers_when_available(self):
        try:
            import openpyxl
        except ImportError:
            self.skipTest('openpyxl not installed')
        wb = openpyxl.load_workbook(BytesIO(self._get(self.user)[1]))
        ws = wb.active
        self.assertEqual(ws['F6'].value, 1200)
        self.assertEqual(ws['F6'].number_format, '#,##0')
        self.assertIsNone(ws['I6'].value)
        self.assertEqual(ws['A8'].value, '金額合計')


class XlsxExportUiTest(unittest.TestCase):
    def _fn(self, name):
        m = re.search(rf'function {name}\([^)]*\)\s*\{{', JS)
        i = JS.find('{', m.start()); d = 0
        for j in range(i, len(JS)):
            d += {'{': 1, '}': -1}.get(JS[j], 0)
            if d == 0:
                return JS[m.start():j + 1]

    def test_button_and_states(self):
        self.assertIn('id="bihinexport"', HTML)
        self.assertIn('エクセル出力', HTML)
        script = self._fn('exportButtonState') + self._fn('filenameFromDisposition') + r'''
console.log(JSON.stringify({
  never: exportButtonState({exists:false}, false), none: exportButtonState(null, false),
  saved: exportButtonState({exists:true}, false), dirty: exportButtonState({exists:true}, true),
  fn: filenameFromDisposition("attachment; filename=\"bihin_2026-10.xlsx\"; filename*=UTF-8''%E5%82%99%E5%93%81%E6%A3%9A%E5%8D%B8_2026-10.xlsx", 'x'),
  fb: filenameFromDisposition(null, 'fallback.xlsx'),
}));'''
        r = subprocess.run(['node', '-e', script], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 0, r.stderr)
        o = json.loads(r.stdout)
        self.assertTrue(o['never']['disabled'])
        self.assertIn('保存されていない', o['never']['title'])
        self.assertTrue(o['none']['disabled'])
        self.assertFalse(o['saved']['disabled'])
        self.assertFalse(o['dirty']['disabled'])
        self.assertIn('最後に保存した内容', o['dirty']['title'])
        self.assertEqual(o['fn'], '備品棚卸_2026-10.xlsx')
        self.assertEqual(o['fb'], 'fallback.xlsx')
        body = self._fn('exportBihinXlsx')
        self.assertIn('confirm(', body)  # 未保存時の警告
        self.assertIn('/api/kintai/bihin/xlsx', body)
        self.assertNotIn('collectPayloadRows', body)  # 画面の未保存値は送らない


if __name__ == '__main__':
    unittest.main()
